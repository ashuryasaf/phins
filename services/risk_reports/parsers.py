"""Risk Reports parsers (B9).

File-format intake for ``AIRiskReportsService``: CSV, Excel (incl. Mislaka
Excel exports), ZIP archives, Mislaka/pension XML, generic XML, images and
PDFs. Every parser returns the same ``{'columns': [...], 'rows': [...]}``
shape the analysis layer consumes.

PDF and image text extraction is delegated to ``DocumentProcessingService``
(pypdf -> regex -> OCR escalation, page-level OCR cache) so both intake paths
share one extractor and one cache.
"""

import copy
import csv
import hashlib
import io
import json
import re
import zipfile
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

from services.risk_reports.models import logger

IMAGE_TYPES = ('png', 'jpg', 'jpeg', 'gif', 'webp')
#: Per-page text kept inline in the rows the analysis layer profiles; the full
#: text is stored once under ``parsed['text']``.
TEXT_ROW_MAX_CHARS = 4000
TEXT_ROW_MAX_PAGES = 50
#: ``DocumentProcessingService`` returns this marker instead of text when a
#: PDF has no readable layer and OCR is unavailable.
NO_TEXT_MARKER_PREFIX = '[PDF content - extraction yielded no text'


def _document_service():
    """Shared Document Intelligence extractor (pypdf -> regex -> OCR, B4 caches)."""
    from services.document_processing_service import get_document_service
    return get_document_service()


class ParserMixin:
    """Format parsers; each returns ``{'columns', 'rows', ...}``."""

    def parse_content(self, filename: str, file_content: bytes,
                      file_type: str) -> Tuple[Dict[str, Any], str]:
        """Dispatch on ``file_type`` and return ``(parsed, encoding)``.

        Pure: nothing is stored, audited or saved here; ``parse_file`` wraps
        this with the document record. Raises on malformed input.
        """
        file_type_lower = (file_type or '').lower()

        if file_type_lower == 'zip':
            return self._parse_zip(file_content), 'utf-8'
        if file_type_lower in IMAGE_TYPES:
            return self._parse_image(file_content, filename, file_type_lower), 'binary'
        if file_type_lower == 'pdf':
            return self._parse_pdf(file_content, filename), 'binary'
        if file_type_lower == 'xml':
            # Mislaka / pension XML (falls back to generic XML inside)
            return self._parse_pension_xml(file_content, filename), 'utf-8'
        if file_type_lower in ('xls', 'xlsx'):
            parsed = self._parse_excel(file_content, filename, file_type_lower)
            if not parsed.get('rows') and not parsed.get('columns'):
                print(f"[AI_REPORTS] Excel parse returned no data for {filename}. "
                      "Ensure openpyxl (xlsx) or xlrd (xls) is installed.")
            return parsed, 'binary'
        if file_type_lower == 'csv':
            encoding = self._detect_encoding(file_content)
            return self._parse_csv(file_content.decode(encoding, errors='replace')), encoding

        # Unknown type: binary -> try Excel, else metadata only; text -> CSV.
        if self._is_binary_content(file_content):
            parsed = self._parse_excel(file_content, filename, 'xlsx')
            if not parsed.get('rows'):
                parsed = {
                    'columns': ['filename', 'size_bytes', 'type', 'note'],
                    'rows': [{
                        'filename': filename,
                        'size_bytes': len(file_content),
                        'type': file_type,
                        'note': 'Binary file - use specific format for full parsing'
                    }],
                    'file_type': 'binary'
                }
            return parsed, 'binary'
        encoding = self._detect_encoding(file_content)
        return self._parse_csv(file_content.decode(encoding, errors='replace')), encoding

    def _detect_encoding(self, content: bytes) -> str:
        """Detect file encoding"""
        # Check for BOM markers
        if content.startswith(b'\xef\xbb\xbf'):
            return 'utf-8-sig'
        if content.startswith(b'\xff\xfe'):
            return 'utf-16-le'
        if content.startswith(b'\xfe\xff'):
            return 'utf-16-be'
        
        # Try common encodings
        for encoding in ['utf-8', 'windows-1255', 'iso-8859-8', 'windows-1252', 'latin-1']:
            try:
                content.decode(encoding)
                return encoding
            except (UnicodeDecodeError, LookupError):
                continue
        
        return 'utf-8'
    
    @staticmethod
    def _is_binary_content(content: bytes) -> bool:
        """
        Detect whether file content is binary (not parseable as plain text).
        Checks for common binary file signatures and control character density.
        """
        if not content:
            return False
        # Check known binary file signatures
        binary_signatures = [
            b'PK',           # ZIP / OOXML (xlsx, docx, etc.)
            b'\xd0\xcf\x11', # OLE2 (xls, doc, etc.)
            b'\x89PNG',      # PNG
            b'\xff\xd8\xff', # JPEG
            b'GIF8',         # GIF
            b'%PDF',         # PDF
            b'RIFF',         # RIFF (webp, wav, etc.)
        ]
        header = content[:8]
        for sig in binary_signatures:
            if header.startswith(sig):
                return True
        # Check for high density of non-text bytes in first 512 bytes
        sample = content[:512]
        non_text = sum(1 for b in sample if b < 0x09 or (0x0E <= b < 0x20 and b != 0x1B))
        return (non_text / max(len(sample), 1)) > 0.10
    
    def _parse_csv(self, text_content: str) -> Dict[str, Any]:
        """Parse CSV content"""
        # Try different delimiters
        for delimiter in [',', ';', '\t', '|']:
            try:
                reader = csv.DictReader(io.StringIO(text_content), delimiter=delimiter)
                rows = list(reader)
                if rows and len(rows[0]) > 1:
                    columns = list(rows[0].keys()) if rows else []
                    return {
                        'columns': columns,
                        'rows': rows,
                        'delimiter': delimiter
                    }
            except Exception:
                continue
        
        # Fallback: simple line parsing
        lines = text_content.strip().split('\n')
        if lines:
            columns = lines[0].split(',')
            rows = []
            for line in lines[1:]:
                values = line.split(',')
                rows.append(dict(zip(columns, values)))
            return {'columns': columns, 'rows': rows, 'delimiter': ','}
        
        return {'columns': [], 'rows': [], 'delimiter': ','}
    
    def _parse_excel(self, content: bytes, filename: str, ext: str) -> Dict[str, Any]:
        """
        Parse Excel files (.xls and .xlsx) from Mislaka data.
        Extracts pension and insurance data from Excel format.
        
        IMPORTANT: XLSX files are ZIP-based OOXML archives. They MUST be parsed
        with openpyxl (xlsx) or xlrd (xls) - never as plain text/CSV, which
        would produce garbled PK!... binary output.
        """
        rows = []
        columns = []
        pension_data = None
        sheet_names = []
        parse_method = None
        
        try:
            if ext == 'xlsx':
                # Use openpyxl for .xlsx files
                try:
                    import openpyxl
                    from openpyxl import load_workbook
                    
                    wb = load_workbook(filename=io.BytesIO(content), data_only=True)
                    sheet_names = wb.sheetnames
                    parse_method = 'openpyxl'
                    
                    for sheet_name in wb.sheetnames:
                        sheet = wb[sheet_name]
                        sheet_rows = list(sheet.iter_rows(values_only=True))
                        
                        if not sheet_rows:
                            continue
                        
                        # First row as headers
                        header_row = sheet_rows[0] if sheet_rows else []
                        sheet_columns = [str(h).strip() if h else f'col_{i}' for i, h in enumerate(header_row)]
                        
                        # Add columns that don't exist yet
                        for col in sheet_columns:
                            if col and col not in columns:
                                columns.append(col)
                        
                        # Process data rows
                        for row in sheet_rows[1:]:
                            if row and any(cell is not None for cell in row):
                                row_dict = {}
                                for i, cell in enumerate(row):
                                    if i < len(sheet_columns):
                                        col_name = sheet_columns[i]
                                        # Convert dates, numbers etc. to string representation
                                        if cell is None:
                                            row_dict[col_name] = ''
                                        elif hasattr(cell, 'isoformat'):
                                            row_dict[col_name] = cell.isoformat()
                                        else:
                                            row_dict[col_name] = cell
                                rows.append(row_dict)
                    
                    wb.close()
                    print(f"[AI_REPORTS] Successfully parsed XLSX '{filename}': "
                          f"{len(sheet_names)} sheets, {len(columns)} columns, {len(rows)} rows")
                    
                except ImportError:
                    print("[AI_REPORTS] WARNING: openpyxl not installed. "
                          "Cannot parse XLSX files. Install with: pip install openpyxl")
                except Exception as e:
                    print(f"[AI_REPORTS] openpyxl error parsing '{filename}': {e}")
                    
            elif ext == 'xls':
                # Use xlrd for older .xls files
                try:
                    import xlrd
                    
                    wb = xlrd.open_workbook(file_contents=content)
                    parse_method = 'xlrd'
                    
                    for sheet_idx in range(wb.nsheets):
                        sheet = wb.sheet_by_index(sheet_idx)
                        sheet_names.append(sheet.name)
                        
                        if sheet.nrows == 0:
                            continue
                        
                        # First row as headers
                        header_row = sheet.row_values(0) if sheet.nrows > 0 else []
                        sheet_columns = [str(h).strip() if h else f'col_{i}' for i, h in enumerate(header_row)]
                        
                        # Add columns that don't exist yet
                        for col in sheet_columns:
                            if col and col not in columns:
                                columns.append(col)
                        
                        # Process data rows
                        for row_idx in range(1, sheet.nrows):
                            row = sheet.row_values(row_idx)
                            if row and any(cell for cell in row):
                                row_dict = {}
                                for i, cell in enumerate(row):
                                    if i < len(sheet_columns):
                                        col_name = sheet_columns[i]
                                        row_dict[col_name] = cell if cell else ''
                                rows.append(row_dict)
                    
                    print(f"[AI_REPORTS] Successfully parsed XLS '{filename}': "
                          f"{len(sheet_names)} sheets, {len(columns)} columns, {len(rows)} rows")
                    
                except ImportError:
                    print("[AI_REPORTS] WARNING: xlrd not installed. "
                          "Cannot parse XLS files. Install with: pip install xlrd")
                except Exception as e:
                    print(f"[AI_REPORTS] xlrd error parsing '{filename}': {e}")
            
            # Try to detect if this is Mislaka pension data
            if rows and columns:
                pension_data = self._detect_mislaka_excel_data(columns, rows, filename)
            
        except Exception as e:
            print(f"[AI_REPORTS] Error parsing Excel file {filename}: {e}")
        
        return {
            'columns': columns,
            'rows': rows,
            'pension_data': pension_data,
            'file_type': 'excel',
            'original_filename': filename,
            'sheet_names': sheet_names,
            'parse_method': parse_method or 'none'
        }
    
    def _detect_mislaka_excel_data(self, columns: List[str], rows: List[Dict], filename: str) -> Optional[Dict[str, Any]]:
        """
        Detect and extract Mislaka pension data from Excel columns/rows.
        Maps Hebrew column names to pension data structure.
        """
        # Hebrew column name mappings for Mislaka data
        column_mappings = {
            # Client fields
            'שם': 'client_name',
            'שם מלא': 'client_name', 
            'שם פרטי': 'first_name',
            'שם משפחה': 'last_name',
            'תעודת זהות': 'id_number',
            'ת.ז': 'id_number',
            'ת"ז': 'id_number',
            'מספר זהות': 'id_number',
            'תאריך לידה': 'birth_date',
            
            # Provider/Product fields
            'יצרן': 'provider',
            'שם יצרן': 'provider',
            'חברה': 'provider',
            'שם חברה': 'provider',
            'מוצר': 'product_name',
            'שם מוצר': 'product_name',
            'סוג מוצר': 'product_type',
            'סוג קופה': 'product_type',
            
            # Policy fields
            'מספר פוליסה': 'policy_number',
            'מס פוליסה': 'policy_number',
            'מספר חשבון': 'policy_number',
            'מס חשבון': 'policy_number',
            
            # Balance fields
            'יתרה': 'balance',
            'יתרה כוללת': 'total_balance',
            'סך צבירה': 'total_balance',
            'צבירה': 'total_balance',
            'סה"כ צבירה': 'total_balance',
            'יתרת תגמולים': 'savings_balance',
            'תגמולים': 'savings_balance',
            'יתרת פיצויים': 'severance_balance',
            'פיצויים': 'severance_balance',
            
            # Fee fields
            'דמי ניהול': 'management_fee',
            'דמי ניהול מצבירה': 'management_fee_savings',
            'דמי ניהול מהפקדות': 'management_fee_deposits',
            'עמלה': 'management_fee',
            
            # Status fields
            'סטטוס': 'status',
            'מצב': 'status',
            'סטטוס פוליסה': 'status',
            
            # Section 14
            'סעיף 14': 'section14',
            'סעיף14': 'section14',
            
            # Employer
            'מעסיק': 'employer_name',
            'שם מעסיק': 'employer_name',
        }
        
        # Check if this looks like pension data
        pension_indicators = ['יצרן', 'פוליסה', 'צבירה', 'יתרה', 'תגמולים', 'פיצויים', 'קופה', 'פנסיה', 'ביטוח', 'גמל']
        columns_lower = [str(c).lower() for c in columns]
        
        is_pension_data = any(
            any(indicator in col for indicator in pension_indicators) 
            for col in columns_lower
        )
        
        if not is_pension_data:
            return None
        
        # Map columns to standardized names
        mapped_columns = {}
        for col in columns:
            col_str = str(col).strip()
            for hebrew_name, english_name in column_mappings.items():
                if hebrew_name in col_str or col_str == hebrew_name:
                    mapped_columns[col] = english_name
                    break
        
        # Extract client and account data
        client_info = {}
        accounts = []
        
        for row in rows:
            account = {}
            
            for original_col, value in row.items():
                if original_col in mapped_columns:
                    mapped_name = mapped_columns[original_col]
                    
                    # Convert value
                    if value is not None and value != '':
                        if mapped_name in ['total_balance', 'savings_balance', 'severance_balance', 
                                          'management_fee', 'management_fee_savings', 'management_fee_deposits']:
                            try:
                                account[mapped_name] = float(str(value).replace(',', '').replace('₪', '').strip())
                            except:
                                account[mapped_name] = 0
                        elif mapped_name == 'section14':
                            account[mapped_name] = str(value).lower() in ['כן', 'yes', '1', 'true', 'v', '✓']
                        elif mapped_name in ['client_name', 'first_name', 'last_name', 'id_number', 'birth_date']:
                            client_info[mapped_name] = str(value).strip()
                        else:
                            account[mapped_name] = str(value).strip()
            
            if account.get('provider') or account.get('policy_number') or account.get('total_balance'):
                accounts.append(account)
        
        # Build full name if we have parts
        if client_info.get('first_name') or client_info.get('last_name'):
            parts = [client_info.get('first_name', ''), client_info.get('last_name', '')]
            client_info['full_name'] = ' '.join(p for p in parts if p)
        elif client_info.get('client_name'):
            client_info['full_name'] = client_info['client_name']

        # Normalize customer identity fields for report consistency.
        client_info = self._normalize_client_profile_fields(client_info)
        
        if not accounts and not client_info:
            return None
        
        # Calculate totals
        total_balance = sum(a.get('total_balance', 0) for a in accounts)
        total_severance = sum(a.get('severance_balance', 0) for a in accounts)
        
        return {
            'client': client_info,
            'accounts': accounts,
            'totals': {
                'total_balance': total_balance,
                'total_balance_formatted': f"₪{total_balance:,.0f}",
                'total_severance': total_severance,
                'total_severance_formatted': f"₪{total_severance:,.0f}",
                'account_count': len(accounts),
                'provider_count': len(set(a.get('provider', '') for a in accounts if a.get('provider'))),
                'providers': list(set(a.get('provider', '') for a in accounts if a.get('provider'))),
            },
            'header': {
                'source': 'Excel',
                'filename': filename,
            }
        }
    
    def _parse_zip(self, content: bytes) -> Dict[str, Any]:
        """Parse ZIP file containing CSV, XML (pension), image, and PDF files"""
        combined_data = {
            'columns': [],
            'rows': [],
            'files': [],
            'pension_data': None,
            'integrity': {
                'zip_file_count': 0,
                'affiliated_files_processed': 0,
                'pension_sources': [],
                'issues': [],
            }
        }
        
        with zipfile.ZipFile(io.BytesIO(content), 'r') as zf:
            for name in zf.namelist():
                # Skip directories and hidden files
                if name.endswith('/') or name.startswith('__') or name.startswith('.'):
                    continue
                combined_data['integrity']['zip_file_count'] += 1
                
                name_lower = name.lower()
                ext = name_lower.split('.')[-1] if '.' in name_lower else ''
                
                with zf.open(name) as f:
                    file_content = f.read()
                
                parsed = None
                file_type = ext
                
                if ext == 'csv':
                    encoding = self._detect_encoding(file_content)
                    text_content = file_content.decode(encoding, errors='replace')
                    parsed = self._parse_csv(text_content)
                elif ext == 'xml':
                    # Check if it's a pension/insurance XML file
                    parsed = self._parse_pension_xml(file_content, name)
                    if parsed:
                        file_type = 'pension_xml'
                        # Store pension data separately for enhanced analysis
                        if parsed.get('pension_data'):
                            combined_data['integrity']['affiliated_files_processed'] += 1
                            combined_data['integrity']['pension_sources'].append(name)
                            combined_data['pension_data'] = self._merge_pension_data_records(
                                combined_data.get('pension_data'),
                                parsed.get('pension_data')
                            )
                elif ext in ['png', 'jpg', 'jpeg', 'gif', 'webp']:
                    parsed = self._parse_image(file_content, name, ext)
                elif ext == 'pdf':
                    parsed = self._parse_pdf(file_content, name)
                elif ext in ['xls', 'xlsx']:
                    # Parse Excel files properly
                    parsed = self._parse_excel(file_content, name, ext)
                    if parsed:
                        file_type = 'excel'
                        # Check if this looks like Mislaka data
                        if parsed.get('pension_data'):
                            combined_data['integrity']['affiliated_files_processed'] += 1
                            combined_data['integrity']['pension_sources'].append(name)
                            combined_data['pension_data'] = self._merge_pension_data_records(
                                combined_data.get('pension_data'),
                                parsed.get('pension_data')
                            )
                
                if parsed:
                    combined_data['files'].append({
                        'name': name,
                        'type': file_type,
                        'columns': parsed.get('columns', []),
                        'row_count': len(parsed.get('rows', []))
                    })
                    
                    # Merge columns and rows
                    for col in parsed.get('columns', []):
                        if col not in combined_data['columns']:
                            combined_data['columns'].append(col)
                    combined_data['rows'].extend(parsed.get('rows', []))

        if (
            combined_data['integrity']['zip_file_count'] > 0
            and combined_data['integrity']['affiliated_files_processed'] == 0
        ):
            combined_data['integrity']['issues'].append(
                'No Swiftness-affiliated XML/Excel records detected in ZIP'
            )
        
        return combined_data

    @staticmethod
    def _normalize_customer_identifier(identifier: Any) -> str:
        """Normalize customer identifiers while preserving non-digit fallback values."""
        text = str(identifier or '').strip()
        if not text:
            return ''

        digits = re.sub(r'\D', '', text)
        # Israeli IDs are 9 digits; some sources drop leading zeroes.
        if 7 <= len(digits) <= 9:
            return digits.zfill(9)
        return text

    @staticmethod
    def _is_valid_israeli_id(identifier: str) -> bool:
        """Validate Israeli ID checksum for 9-digit identifiers."""
        if not identifier or not identifier.isdigit() or len(identifier) != 9:
            return False

        total = 0
        for index, char in enumerate(identifier):
            digit = int(char)
            factor = 1 if index % 2 == 0 else 2
            product = digit * factor
            if product > 9:
                product -= 9
            total += product

        return total % 10 == 0

    @staticmethod
    def _normalize_birth_date(value: Any) -> Tuple[str, str]:
        """
        Normalize birth dates into:
          - raw canonical format: YYYYMMDD
          - display format: DD/MM/YYYY
        """
        raw_input = str(value or '').strip()
        if not raw_input:
            return '', ''

        def _to_pair(dt: datetime) -> Tuple[str, str]:
            return dt.strftime('%Y%m%d'), dt.strftime('%d/%m/%Y')

        digits = re.sub(r'\D', '', raw_input)
        if len(digits) == 8:
            # Prefer YYYYMMDD (e.g. 19781111), fallback to DDMMYYYY.
            for year, month, day in [
                (digits[0:4], digits[4:6], digits[6:8]),
                (digits[4:8], digits[2:4], digits[0:2]),
            ]:
                try:
                    parsed = datetime(int(year), int(month), int(day))
                    if 1900 <= parsed.year <= 2100:
                        return _to_pair(parsed)
                except Exception:
                    continue

        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%Y/%m/%d', '%d-%m-%Y', '%Y.%m.%d', '%d.%m.%Y'):
            try:
                parsed = datetime.strptime(raw_input, fmt)
                if 1900 <= parsed.year <= 2100:
                    return _to_pair(parsed)
            except Exception:
                continue

        try:
            parsed = datetime.fromisoformat(raw_input.replace('Z', '+00:00'))
            if 1900 <= parsed.year <= 2100:
                return _to_pair(parsed)
        except Exception:
            pass

        return digits if len(digits) == 8 else raw_input, ''

    def _normalize_client_profile_fields(self, client_info: Optional[Dict[str, Any]]) -> Dict[str, Any]:
        """Normalize and enrich client identity fields used in reports."""
        normalized = dict(client_info or {})

        normalized_id = self._normalize_customer_identifier(normalized.get('id_number'))
        if normalized_id:
            normalized['id_number'] = normalized_id
            normalized['id_israeli_valid'] = self._is_valid_israeli_id(normalized_id)

        birth_raw, birth_display = self._normalize_birth_date(normalized.get('birth_date'))
        if birth_raw:
            normalized['birth_date_raw'] = birth_raw
        if birth_display:
            normalized['birth_date'] = birth_display

        return normalized

    def _merge_pension_data_records(
        self,
        current: Optional[Dict[str, Any]],
        incoming: Optional[Dict[str, Any]]
    ) -> Dict[str, Any]:
        """
        Merge multiple pension-affiliated payloads from ZIP members while preserving
        customer profile and totals integrity.
        """
        if not isinstance(current, dict) or not current:
            base = copy.deepcopy(incoming or {})
            if isinstance(base.get('client'), dict):
                base['client'] = self._normalize_client_profile_fields(base.get('client', {}))
            return base
        if not isinstance(incoming, dict) or not incoming:
            return copy.deepcopy(current)

        merged = copy.deepcopy(current)
        incoming_copy = copy.deepcopy(incoming)

        def _dedupe_by_key(rows: List[Dict[str, Any]], key_fields: List[str]) -> List[Dict[str, Any]]:
            results: List[Dict[str, Any]] = []
            seen: set = set()
            for row in rows:
                if not isinstance(row, dict):
                    continue
                key = tuple(str(row.get(field, '') or '').strip() for field in key_fields)
                if not any(key):
                    payload = json.dumps(row, sort_keys=True, ensure_ascii=False)
                    key = ('hash', hashlib.sha256(payload.encode('utf-8')).hexdigest())
                if key in seen:
                    continue
                seen.add(key)
                results.append(row)
            return results

        merged_client = merged.get('client', {})
        incoming_client = incoming_copy.get('client', {})
        if isinstance(merged_client, list):
            merged_client = merged_client[0] if merged_client else {}
        if isinstance(incoming_client, list):
            incoming_client = incoming_client[0] if incoming_client else {}

        merged_client = self._normalize_client_profile_fields(merged_client if isinstance(merged_client, dict) else {})
        incoming_client = self._normalize_client_profile_fields(incoming_client if isinstance(incoming_client, dict) else {})
        for key, value in incoming_client.items():
            if value and not merged_client.get(key):
                merged_client[key] = value
        merged_anomalies = list(merged.get('anomalies', []) or [])
        if (
            merged_client.get('id_number')
            and incoming_client.get('id_number')
            and str(merged_client.get('id_number')) != str(incoming_client.get('id_number'))
        ):
            merged_anomalies.append(
                f"Client ID mismatch across affiliated files: {merged_client.get('id_number')} vs {incoming_client.get('id_number')}"
            )
        merged['anomalies'] = merged_anomalies
        merged['client'] = merged_client

        merged_accounts = list(merged.get('accounts', []) or [])
        incoming_accounts = list(incoming_copy.get('accounts', []) or [])
        merged['accounts'] = _dedupe_by_key(
            merged_accounts + incoming_accounts,
            ['policy_number', 'provider', 'product_type', 'start_date']
        )

        merged_contributions = list(merged.get('contributions', []) or [])
        incoming_contributions = list(incoming_copy.get('contributions', []) or [])
        merged['contributions'] = _dedupe_by_key(
            merged_contributions + incoming_contributions,
            ['period', 'policy_number', 'employer_name', 'employee_amount', 'employer_amount', 'severance_amount', 'total_amount']
        )

        merged_severance = list(merged.get('severance', []) or [])
        incoming_severance = list(incoming_copy.get('severance', []) or [])
        merged['severance'] = _dedupe_by_key(
            merged_severance + incoming_severance,
            ['employer_name', 'section14_date', 'total_severance']
        )

        merged_employers = list(merged.get('employers', []) or [])
        incoming_employers = list(incoming_copy.get('employers', []) or [])
        employer_rows: List[Dict[str, Any]] = []
        for employer in (merged_employers + incoming_employers):
            if isinstance(employer, dict):
                employer_rows.append(employer)
            else:
                name = str(employer or '').strip()
                if name:
                    employer_rows.append({'id': '', 'name': name})
        merged['employers'] = _dedupe_by_key(employer_rows, ['id', 'name'])

        merged_providers = list(merged.get('providers', []) or [])
        incoming_providers = list(incoming_copy.get('providers', []) or [])
        if merged_providers and isinstance(merged_providers[0], dict):
            merged['providers'] = _dedupe_by_key(merged_providers + incoming_providers, ['code', 'name'])
        else:
            merged['providers'] = sorted({
                str(p).strip()
                for p in (merged_providers + incoming_providers)
                if str(p).strip()
            })

        header = dict(merged.get('header', {}) or {})
        incoming_header = dict(incoming_copy.get('header', {}) or {})
        for key, value in incoming_header.items():
            if value and not header.get(key):
                header[key] = value
        merged['header'] = header

        # Recompute key totals after merge.
        total_balance = sum(self._to_float_amount(a.get('total_balance')) for a in merged.get('accounts', []))
        total_savings = sum(self._to_float_amount(a.get('savings_balance')) for a in merged.get('accounts', []))
        total_severance = (
            sum(self._to_float_amount(a.get('severance_balance')) for a in merged.get('accounts', [])) +
            sum(self._to_float_amount(s.get('total_severance')) for s in merged.get('severance', []))
        )
        provider_names = sorted({
            str(a.get('provider', '')).strip()
            for a in merged.get('accounts', [])
            if str(a.get('provider', '')).strip()
        })
        totals = dict(merged.get('totals', {}) or {})
        totals.update({
            'total_balance': round(total_balance, 2),
            'total_balance_formatted': f"₪{total_balance:,.2f}",
            'total_savings': round(total_savings, 2),
            'total_savings_formatted': f"₪{total_savings:,.2f}",
            'total_severance': round(total_severance, 2),
            'total_severance_formatted': f"₪{total_severance:,.2f}",
            'total_coverage': round(
                sum(
                    self._to_float_amount(a.get('coverage_amount'))
                    or self._to_float_amount(a.get('death_coverage')) + self._to_float_amount(a.get('disability_coverage'))
                    for a in merged.get('accounts', [])
                ),
                2
            ),
            'account_count': len(merged.get('accounts', [])),
            'provider_count': len(provider_names),
            'providers': provider_names,
            'section14_coverage': any(bool(a.get('section14')) for a in merged.get('accounts', [])),
        })
        merged['totals'] = totals

        return merged
    
    def _parse_pension_xml(self, content: bytes, filename: str) -> Dict[str, Any]:
        """
        Parse Israeli pension/insurance XML files using the PensionDataAgent.
        
        Supports Mislaka (מסלקה) interface standards:
        - Type 1: Holdings (אחזקות)
        - Type 2: Pre-Advice (הודעה מקדימה)
        - Type 3: Holdings + Pre-Advice Combined
        - Type 17: Severance (פיצויים)
        """
        try:
            from services.pension_data_agent import get_pension_agent, is_pension_xml
            
            # Check if this is a pension XML file
            if not is_pension_xml(content):
                # Not a pension XML, try to parse as generic XML
                return self._parse_generic_xml(content, filename)
            
            # Process with PensionDataAgent
            agent = get_pension_agent()
            result = agent.process_xml_content(content)
            
            pension_data = result.get('data', {})
            if isinstance(pension_data.get('client'), dict):
                pension_data['client'] = self._normalize_client_profile_fields(pension_data.get('client', {}))
            report_text = result.get('report', '')
            
            # Convert to CSV-like format for AI analysis
            columns, rows = agent.to_csv_format(pension_data)
            
            # Add summary data as additional rows
            summary = pension_data.get('summary', {})
            header = pension_data.get('header', {})
            clients = pension_data.get('client', [])
            
            # Add header info as rows
            meta_rows = [
                {'מספר פוליסה': 'סוג ממשק', 'יצרן': pension_data.get('interface_type', ''), 'סוג מוצר': '', 'שם מוצר': '', 'סטטוס': '', 'יתרה': '', 'פיצויים': '', 'מעסיק': ''},
                {'מספר פוליסה': 'גרסת סכמה', 'יצרן': header.get('schema_version', ''), 'סוג מוצר': '', 'שם מוצר': '', 'סטטוס': '', 'יתרה': '', 'פיצויים': '', 'מעסיק': ''},
            ]
            
            # Add client info
            if clients:
                client = clients[0] if isinstance(clients, list) else clients
                meta_rows.append({
                    'מספר פוליסה': 'לקוח',
                    'יצרן': client.get('name', ''),
                    'סוג מוצר': '',
                    'שם מוצר': '',
                    'סטטוס': '',
                    'יתרה': '',
                    'פיצויים': '',
                    'מעסיק': ''
                })
            
            # Prepend meta rows
            all_rows = meta_rows + rows
            
            return {
                'columns': columns,
                'rows': all_rows,
                'delimiter': None,
                'file_type': 'pension_xml',
                'original_filename': filename,
                'pension_data': pension_data,
                'pension_report': report_text
            }
            
        except ImportError:
            print("[AI_REPORTS] PensionDataAgent not available, falling back to generic XML parsing")
            return self._parse_generic_xml(content, filename)
        except Exception as e:
            print(f"[AI_REPORTS] Error parsing pension XML: {e}")
            return self._parse_generic_xml(content, filename)
    
    def _parse_generic_xml(self, content: bytes, filename: str) -> Dict[str, Any]:
        """
        Parse generic XML file into tabular format.
        """
        try:
            from defusedxml import ElementTree as DefusedET
            
            # Decode content
            encoding = self._detect_encoding(content)
            xml_str = content.decode(encoding, errors='replace')
            
            # Parse XML
            root = DefusedET.fromstring(xml_str)
            
            # Extract all leaf elements as rows
            columns = ['element', 'value', 'path']
            rows = []
            
            def extract_elements(elem, path=""):
                current_path = f"{path}/{elem.tag}" if path else elem.tag
                
                # If element has text content
                if elem.text and elem.text.strip():
                    rows.append({
                        'element': elem.tag,
                        'value': elem.text.strip(),
                        'path': current_path
                    })
                
                # Process attributes
                for attr, val in elem.attrib.items():
                    rows.append({
                        'element': f"{elem.tag}@{attr}",
                        'value': val,
                        'path': current_path
                    })
                
                # Process children
                for child in elem:
                    extract_elements(child, current_path)
            
            extract_elements(root)
            
            return {
                'columns': columns,
                'rows': rows,
                'delimiter': None,
                'file_type': 'xml',
                'original_filename': filename
            }
            
        except Exception as e:
            print(f"[AI_REPORTS] Error parsing generic XML: {e}")
            return {'columns': [], 'rows': [], 'file_type': 'xml', 'error': str(e)}
    

    # -- images & PDFs (delegate text extraction to Document Intelligence) --
    def _extractor(self):
        """The shared extractor, or ``None`` when Document Intelligence is unavailable."""
        try:
            return _document_service()
        except Exception as exc:  # pragma: no cover - import/boot failure only
            logger.warning("risk reports: document extractor unavailable (%s)", exc)
            return None

    @staticmethod
    def _text_rows(text: str, pages: List[Dict[str, int]]) -> List[Dict[str, Any]]:
        """One ``content`` row per page (or one for the whole text) for the analysis layer."""
        if not text:
            return []
        if pages:
            chunks = [(p.get('page', i + 1), text[p['char_start']:p['char_end']])
                      for i, p in enumerate(pages[:TEXT_ROW_MAX_PAGES])]
        else:
            chunks = [(1, text)]
        rows = []
        for page_no, chunk in chunks:
            chunk = (chunk or '').strip()
            if chunk:
                rows.append({'property': f'page_{page_no}_text',
                             'value': chunk[:TEXT_ROW_MAX_CHARS], 'category': 'content'})
        return rows

    def _parse_image(self, content: bytes, filename: str, file_type: str) -> Dict[str, Any]:
        """
        Parse an image: dimensions plus any OCR text.

        Dimensions and OCR come from ``DocumentProcessingService`` (Tesseract
        when installed, page-level cache). Without OCR the result is the
        historical metadata-only table.
        """
        file_size = len(content)
        extractor = self._extractor()

        width, height = 0, 0
        ocr_text = ''
        if extractor is not None:
            try:
                info = extractor._image_metadata(content)
                width = int(info.get('width') or 0)
                height = int(info.get('height') or 0)
            except Exception as exc:
                logger.debug("image metadata failed: %s", exc)
            try:
                ocr_text = extractor._ocr_image_bytes(
                    content, lang_hint=self._filename_lang_hint(filename)) or ''
            except Exception as exc:
                logger.debug("image OCR failed: %s", exc)
                ocr_text = ''

        # Extract any text from filename for language detection
        filename_text = filename.replace('_', ' ').replace('-', ' ')

        columns = ['property', 'value', 'category']
        rows = [
            {'property': 'filename', 'value': filename, 'category': 'metadata'},
            {'property': 'file_type', 'value': file_type.upper(), 'category': 'metadata'},
            {'property': 'file_size_bytes', 'value': str(file_size), 'category': 'metadata'},
            {'property': 'file_size_kb', 'value': str(round(file_size / 1024, 2)), 'category': 'metadata'},
            {'property': 'width_px', 'value': str(width), 'category': 'dimensions'},
            {'property': 'height_px', 'value': str(height), 'category': 'dimensions'},
            {'property': 'resolution', 'value': f'{width}x{height}', 'category': 'dimensions'},
            {'property': 'document_type', 'value': 'image', 'category': 'classification'},
            {'property': 'source_name', 'value': filename_text, 'category': 'context'},
        ]
        ocr_text = ocr_text.strip()
        if ocr_text:
            rows.append({'property': 'ocr_text', 'value': ocr_text[:TEXT_ROW_MAX_CHARS],
                         'category': 'content'})

        parsed: Dict[str, Any] = {
            'columns': columns,
            'rows': rows,
            'delimiter': None,
            'file_type': 'image',
            'original_filename': filename,
        }
        if ocr_text:
            parsed['text'] = ocr_text
            parsed['text_extraction'] = 'ocr'
        return parsed

    @staticmethod
    def _filename_lang_hint(filename: str) -> Optional[str]:
        return 'hebrew' if re.search(r'[\u0590-\u05FF]', filename or '') else None

    def _parse_pdf(self, content: bytes, filename: str) -> Dict[str, Any]:
        """
        Parse a PDF: metadata plus the extracted text, one row per page.

        Text comes from ``DocumentProcessingService`` (pypdf text layer,
        regex fallback, then OCR for scanned files) so Hebrew policy fields in
        the document itself - not only in its filename - reach the analysis.
        """
        file_size = len(content)

        # Basic PDF metadata extraction
        page_count = 0
        title = filename
        author = ''

        # Simple PDF parsing for metadata
        try:
            content_str = content[:4096].decode('latin-1', errors='ignore')

            # Count pages (approximate; ``/Type /Pages`` nodes are excluded)
            page_count = len(re.findall(rb'/Type\s*/Page(?!s)', content))

            # Extract title if present
            if '/Title' in content_str:
                start = content_str.find('/Title')
                if start != -1:
                    # Try to extract title value
                    paren_start = content_str.find('(', start)
                    paren_end = content_str.find(')', paren_start) if paren_start != -1 else -1
                    if paren_start != -1 and paren_end != -1:
                        title = content_str[paren_start+1:paren_end][:100]

            # Extract author if present
            if '/Author' in content_str:
                start = content_str.find('/Author')
                if start != -1:
                    paren_start = content_str.find('(', start)
                    paren_end = content_str.find(')', paren_start) if paren_start != -1 else -1
                    if paren_start != -1 and paren_end != -1:
                        author = content_str[paren_start+1:paren_end][:100]
        except Exception:
            pass

        # Extract text from filename for context
        filename_text = filename.replace('_', ' ').replace('-', ' ').replace('.pdf', '')

        text, pages = '', []
        extractor = self._extractor()
        if extractor is not None:
            try:
                text, pages = extractor._extract_pdf_text_with_pages(
                    content, lang_hint=self._filename_lang_hint(filename))
                # The extractor returns an explanatory marker, not text, when
                # nothing could be read; keep the table metadata-only then.
                if text.startswith(NO_TEXT_MARKER_PREFIX) or not extractor._has_meaningful_text(text):
                    text, pages = '', []
                elif pages:
                    # The extractor read the page tree; it caps very long
                    # documents, so never lower a larger structural count.
                    page_count = max(page_count, len(pages))
            except Exception as exc:
                logger.debug("pdf text extraction failed: %s", exc)
                text, pages = '', []

        # Detect Hebrew in the filename or the extracted text
        has_hebrew = bool(re.search(r'[\u0590-\u05FF]', filename_text)) or \
            bool(re.search(r'[\u0590-\u05FF]', text[:20000]))

        columns = ['property', 'value', 'category']
        rows = [
            {'property': 'filename', 'value': filename, 'category': 'metadata'},
            {'property': 'file_type', 'value': 'PDF', 'category': 'metadata'},
            {'property': 'file_size_bytes', 'value': str(file_size), 'category': 'metadata'},
            {'property': 'file_size_kb', 'value': str(round(file_size / 1024, 2)), 'category': 'metadata'},
            {'property': 'file_size_mb', 'value': str(round(file_size / (1024*1024), 2)), 'category': 'metadata'},
            {'property': 'page_count', 'value': str(page_count), 'category': 'content'},
            {'property': 'title', 'value': title, 'category': 'metadata'},
            {'property': 'author', 'value': author, 'category': 'metadata'},
            {'property': 'document_type', 'value': 'pdf', 'category': 'classification'},
            {'property': 'source_name', 'value': filename_text, 'category': 'context'},
            {'property': 'has_hebrew', 'value': str(has_hebrew), 'category': 'language'},
        ]
        text_rows = self._text_rows(text, pages)
        rows.extend(text_rows)

        parsed: Dict[str, Any] = {
            'columns': columns,
            'rows': rows,
            'delimiter': None,
            'file_type': 'pdf',
            'original_filename': filename,
            'page_count': page_count,
        }
        if text_rows:
            parsed['text'] = text
            parsed['text_pages'] = pages
            parsed['text_extraction'] = 'document_intelligence'
        return parsed
