"""Mislaka file parsers: XML (tree and streaming), Excel and CSV.

Moved from ``services/pension_data_agent.py`` (B5) as a mixin. Behavioural
changes are confined to *how* the XML is walked, never to the output shape:

* ``_find_text`` reads a per-parse descendant index instead of running up to
  three ``.//tag`` scans per field (the index records the first element per
  tag in document order, exactly what ``find('.//tag')`` returned).
* Files larger than ``PHINS_PENSION_STREAM_MIN_BYTES`` are parsed with
  ``defusedxml`` ``iterparse``: each provider block (``YeshutYatzran``) is
  parsed and released as soon as it closes, so peak memory is bounded by the
  largest provider block rather than the whole document.
"""

import io
import logging
import os
import re
import threading
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

from defusedxml import ElementTree as defused_etree

from services.pension.schema import (
    CompiledFields,
    MislakaSchemaMapping,
    looks_like_pension_table,
    map_hebrew_column,
    tag_variants,
)

logger = logging.getLogger('services.pension_data_agent')


class MislakaParserMixin:
    """XML / Excel / CSV → the parser's canonical ``data`` dict."""

    schema_mapping: MislakaSchemaMapping

    # ------------------------------------------------------------------------
    # XML entry points
    # ------------------------------------------------------------------------

    STREAM_MIN_BYTES_ENV = 'PHINS_PENSION_STREAM_MIN_BYTES'
    DEFAULT_STREAM_MIN_BYTES = 8 * 1024 * 1024

    PROVIDER_TAG = 'YeshutYatzran'
    HEADER_TAGS = ('KoteretKovetz', 'Header')
    CLIENT_TAGS = ('YeshutLakoach', 'YeshutLakohach', 'Lakoach', 'Mevutach', 'Client', 'ClientDetails')
    PRODUCT_TAG = 'Mutzar'
    ACCOUNT_TAGS = ('HeshbonOPolisa', 'PirteiHeshbon', 'Account', 'Policy', 'Plan', 'ReshimatKupa')
    STANDALONE_ACCOUNT_TAGS = ('HeshbonOPolisa', 'Account', 'Policy')
    CONTRIBUTION_TAGS = ('NetuneiHafrasha', 'PirteiHafrasha', 'Hafrasha', 'ReshimatHafrashot', 'Peula', 'Event', 'Transaction')
    SEVERANCE_TAGS = ('NetuneiPitzuim', 'PirteiPitzuim', 'Pitzuim', 'SeveranceDetails')
    EMPLOYER_TAG = 'YeshutMaasik'
    YITRA_BLOCK_TAGS = (
        'Yitra', 'PerutYitra', 'PerutYitraLeFiSugHafrasha',
        'YitraLefiSugHafrasha', 'YitraLeFiSugHafrasha',
    )
    YITRA_AMOUNT_TAGS = (
        'SCHUM-TZVIRA', 'SchumTzvira', 'SCHUM-TZVIRA-NOCHECHIT',
        'TOTAL-CHISACHON', 'TotalChisachon', 'TOTAL-CHISACHON-MTZBR',
        'ERECH-PIDYON', 'ErechPidyon', 'ERECH-PIDYON-NOCHECHI',
        'SACH-YITRA', 'SachYitra', 'SCHUM', 'Saldo', 'SALDO',
    )
    YITRA_TYPE_TAGS = ('KOD-SUG-HAFRASHA', 'KodSugHafrasha', 'SUG-YITRA', 'SUG-HAFRASHA')
    TOTAL_BALANCE_VARIANTS = tuple(
        variants for _tag, field, variants in CompiledFields.ACCOUNT if field == 'total_balance'
    )
    CLIENT_ID_RAW_TAGS = (
        'MISPAR-ZIHUI-LAKOACH', 'MisparZihuiLakoach', 'MISPARZEHUT',
        'MisparZehut', 'MISPAR-ZEHUT', 'MISPAR-ZIHUY', 'ZEHUT', 'TEUDAT-ZEHUT',
    )

    _parse_local = threading.local()

    @staticmethod
    def _local_tag(tag: Optional[str]) -> str:
        """Strip Clark-notation namespaces used by official Swiftness/Mislaka XML."""
        if not tag:
            return ''
        if tag[0] == '{' and '}' in tag:
            return tag.rsplit('}', 1)[-1]
        return tag

    def _descendants_named(self, elem, *names, include_self: bool = False):
        """Yield descendants whose local tag is in ``names`` (namespace-safe)."""
        wanted = set(names)
        skip_self = not include_self
        for node in elem.iter():
            if skip_self:
                skip_self = False
                continue
            if self._local_tag(node.tag) in wanted:
                yield node

    def _first_named(self, elem, *names, include_self: bool = False):
        for node in self._descendants_named(elem, *names, include_self=include_self):
            return node
        return None

    @classmethod
    def stream_min_bytes(cls) -> int:
        """Documents at least this large are parsed with ``iterparse``."""
        try:
            return max(0, int(os.environ.get(cls.STREAM_MIN_BYTES_ENV, cls.DEFAULT_STREAM_MIN_BYTES)))
        except (TypeError, ValueError):
            return cls.DEFAULT_STREAM_MIN_BYTES

    def _parse_mislaka_xml(self, xml_content: bytes) -> Dict[str, Any]:
        """Parse Mislaka XML into structured data with proper encoding handling.

        Large documents take the streaming path; anything it cannot handle
        (an undeclared legacy encoding, malformed markup) falls back to the
        tree parser, which owns the encoding fallbacks and the final error.
        """
        if len(xml_content) >= self.stream_min_bytes():
            try:
                return self._parse_mislaka_xml_streaming(xml_content)
            except Exception as exc:  # fall through to the tree parser, which decides
                logger.warning("streaming Mislaka parse failed (%s); using tree parse", exc)
        root = self._load_root(xml_content)
        with self._parse_context():
            return self._parse_root(root)

    def _load_root(self, xml_content: bytes):
        # Try to parse with automatic encoding detection
        try:
            try:
                root = defused_etree.fromstring(xml_content)
            except Exception:
                try:
                    root = defused_etree.fromstring(xml_content.decode('utf-8', errors='replace'))
                except Exception:
                    root = defused_etree.fromstring(xml_content.decode('windows-1255', errors='replace'))
        except Exception as e:
            # Last resort: try Windows-1255
            try:
                text = xml_content.decode('windows-1255', errors='replace')
                root = defused_etree.fromstring(text)
            except Exception as e2:
                raise ValueError(f"Failed to parse XML: {e2}")
        return root

    def _parse_root(self, root) -> Dict[str, Any]:
        """Tree parse: the historical algorithm, unchanged in output."""
        data = self._empty_data()

        # Parse header
        data['header'] = self._parse_header(root)
        self._set_interface(data)

        # Parse client
        data['client'] = self._parse_client(root)

        # Parse providers and accounts
        providers, accounts = self._parse_providers_and_accounts(root)
        data['providers'] = providers
        data['accounts'] = accounts

        # Parse contributions
        data['contributions'] = self._parse_contributions(root)

        # Parse severance data
        data['severance'] = self._parse_severance(root, data['interface_code'])

        # Parse employers
        data['employers'] = self._parse_employers(root, accounts)

        # Extract all raw text elements for additional analysis
        data['raw_elements'] = self._extract_all_elements(root)
        self._recover_client_identity(data)

        return data

    @staticmethod
    def _empty_data() -> Dict[str, Any]:
        return {
            'header': {},
            'client': {},
            'providers': [],
            'accounts': [],
            'contributions': [],
            'severance': [],
            'employers': [],
            'raw_elements': {},
        }

    def _set_interface(self, data: Dict[str, Any]) -> None:
        interface_code = data['header'].get('interface_code', 1)
        data['interface_code'] = interface_code
        interface_info = self.schema_mapping.INTERFACE_CODES.get(
            interface_code,
            {'name': f'Type{interface_code}', 'he': f'סוג {interface_code}'}
        )
        data['interface_type'] = interface_info['name']
        data['interface_type_he'] = interface_info['he']

    # ------------------------------------------------------------------------
    # Streaming parse (B5)
    # ------------------------------------------------------------------------

    def _parse_mislaka_xml_streaming(self, xml_content: bytes) -> Dict[str, Any]:
        """``iterparse`` variant of ``_parse_root``.

        Every top-level provider block is parsed when its end tag arrives and
        then removed from the tree, so the document never has to be resident
        in full. Candidates that the tree parser located with a document-wide
        ``.//`` search (header, client, standalone accounts, contributions,
        severance, employers) are collected per block with their document
        position and assembled afterwards in the tree parser's order.
        ``raw_elements`` is filled from the end events themselves.

        Known difference from the tree parser, by construction: when a
        document has no header element at all, the header-field fallback
        searches only the part of the document outside provider blocks.
        """
        acc = _StreamAccumulator()
        stack: List[Any] = []
        provider_depth = 0
        root = None
        position = 0

        with self._parse_context() as ctx:
            for event, elem in defused_etree.iterparse(io.BytesIO(xml_content), events=('start', 'end')):
                if event == 'start':
                    if root is None:
                        root = elem
                    stack.append(elem)
                    if self._local_tag(elem.tag) == self.PROVIDER_TAG:
                        provider_depth += 1
                    continue

                # -- end event -------------------------------------------------
                position += 1
                stack.pop()
                tag = self._local_tag(elem.tag)
                text = elem.text
                if text and text.strip():
                    acc.raw.setdefault(tag, []).append(text.strip())

                if tag in acc.CANDIDATE_TAGS:
                    ctx.positions[id(elem)] = position
                if tag == self.PROVIDER_TAG:
                    provider_depth -= 1
                    if provider_depth == 0:
                        self._harvest_block(elem, acc, ctx, is_provider_block=True)
                        self._release_block(elem, stack[-1] if stack else None, ctx)

            if root is None:
                raise ValueError("Failed to parse XML: empty document")

            # Residual tree: everything that was not inside a provider block.
            self._harvest_block(root, acc, ctx, is_provider_block=False)
            return self._assemble(root, acc, ctx)

    def _harvest_block(self, block, acc: '_StreamAccumulator', ctx: '_ParseContext',
                       is_provider_block: bool) -> None:
        """Parse one subtree now, while it is still resident."""
        positions = ctx.positions

        def descendants(tag):
            return list(self._descendants_named(block, tag))

        def pos_of(elem):
            return positions.get(id(elem), float('inf'))

        for tag in self.HEADER_TAGS:
            for elem in descendants(tag):
                acc.header_candidates.append((self.HEADER_TAGS.index(tag), pos_of(elem), self._header_from_elem(elem)))
                break
        for tag in self.CLIENT_TAGS:
            for elem in descendants(tag):
                acc.client_candidates.append((self.CLIENT_TAGS.index(tag), pos_of(elem), self._client_from_elem(elem)))
                break

        product_account_ids: set = set()
        provider_elems = ([block] if is_provider_block else []) + descendants(self.PROVIDER_TAG)
        for provider_elem in provider_elems:
            self._parse_provider_elem(provider_elem, acc.providers, acc.provider_accounts, product_account_ids)

        for tag in self.STANDALONE_ACCOUNT_TAGS:
            for elem in descendants(tag):
                if id(elem) in product_account_ids:
                    continue  # already parsed under a product; its policy number is on file
                policy_num = self._find_text(elem, 'MISPAR-POLISA-O-HESHBON')
                if policy_num:
                    acc.standalone.append((self.STANDALONE_ACCOUNT_TAGS.index(tag), pos_of(elem),
                                           policy_num, self._parse_account(elem, {}, {})))

        for tag in self.CONTRIBUTION_TAGS:
            for elem in descendants(tag):
                contrib = self._contribution_from_elem(elem)
                if contrib:
                    acc.contributions.append((self.CONTRIBUTION_TAGS.index(tag), pos_of(elem), contrib))

        for tag in self.SEVERANCE_TAGS:
            for elem in descendants(tag):
                sev = self._severance_from_elem(elem)
                if sev:
                    acc.severance.append((self.SEVERANCE_TAGS.index(tag), pos_of(elem), sev))

        for elem in descendants(self.EMPLOYER_TAG):
            emp_name = self._find_text(elem, 'SHEM-MAASIK') or self._find_text(elem, 'ShemMaasik')
            emp_id = self._find_text(elem, 'KOD-MAASIK') or self._find_text(elem, 'KodMaasik')
            acc.employer_elems.append((pos_of(elem), emp_name, emp_id))

    def _release_block(self, elem, parent, ctx: '_ParseContext') -> None:
        """Drop a parsed provider block and every per-parse note about it, so a
        recycled ``id()`` can never alias a stale index or position."""
        for d in elem.iter():
            ctx.index.pop(id(d), None)
            ctx.positions.pop(id(d), None)
        elem.clear()
        if parent is not None:
            try:
                parent.remove(elem)
            except ValueError:
                pass

    def _assemble(self, root, acc: '_StreamAccumulator', ctx: '_ParseContext') -> Dict[str, Any]:
        data = self._empty_data()

        if acc.header_candidates:
            data['header'] = min(acc.header_candidates, key=lambda c: (c[0], c[1]))[2]
        else:
            data['header'] = self._header_from_elem(root)
        self._set_interface(data)

        if acc.client_candidates:
            data['client'] = min(acc.client_candidates, key=lambda c: (c[0], c[1]))[2]

        accounts = list(acc.provider_accounts)
        for _, _, policy_num, account in sorted(acc.standalone, key=lambda c: (c[0], c[1])):
            if account and not any(a.get('policy_number') == policy_num for a in accounts):
                accounts.append(account)
        data['providers'] = acc.providers
        data['accounts'] = accounts
        data['contributions'] = [c[2] for c in sorted(acc.contributions, key=lambda c: (c[0], c[1]))]
        data['severance'] = [s[2] for s in sorted(acc.severance, key=lambda c: (c[0], c[1]))]

        employers: List[Dict[str, Any]] = []
        seen: set = set()
        for acct in accounts:
            emp_id = acct.get('employer_id', '')
            emp_name = acct.get('employer_name', '')
            if emp_name and emp_name not in seen:
                seen.add(emp_name)
                employers.append({'id': emp_id, 'name': emp_name})
        for _, emp_name, emp_id in sorted(acc.employer_elems, key=lambda c: c[0]):
            if emp_name and emp_name not in seen:
                seen.add(emp_name)
                employers.append({'id': emp_id, 'name': emp_name})
        data['employers'] = employers
        data['raw_elements'] = acc.raw
        self._recover_client_identity(data)
        return data

    # ------------------------------------------------------------------------
    # Per-parse context: descendant text index (replaces repeated .// scans)
    # ------------------------------------------------------------------------

    class _parse_context:
        """Context manager binding a fresh ``_ParseContext`` to this thread."""

        def __init__(self):
            self.ctx = _ParseContext()
            self.previous = None

        def __enter__(self):
            local = MislakaParserMixin._parse_local
            self.previous = getattr(local, 'ctx', None)
            local.ctx = self.ctx
            return self.ctx

        def __exit__(self, *exc):
            MislakaParserMixin._parse_local.ctx = self.previous
            return False

    def _forget_index(self, elem) -> None:
        """Drop the descendant indexes built under ``elem`` (parsed, done)."""
        ctx = getattr(self._parse_local, 'ctx', None)
        if ctx is None or not ctx.index:
            return
        for d in elem.iter():
            ctx.index.pop(id(d), None)

    def _text_index(self, elem) -> Dict[str, Optional[str]]:
        """``tag → text of the first descendant with that tag`` (document
        order), which is exactly what ``elem.find('.//tag')`` resolved to."""
        ctx = getattr(self._parse_local, 'ctx', None)
        if ctx is None:
            return None
        index = ctx.index.get(id(elem))
        if index is None:
            index = {}
            first = True
            for d in elem.iter():
                if first:
                    first = False
                    continue
                local = self._local_tag(d.tag)
                if local not in index:
                    index[local] = d.text
            ctx.index[id(elem)] = index
        return index

    def _parse_mislaka_excel(self, content: bytes, filename: str, ext: str) -> Optional[Dict[str, Any]]:
        """
        Parse Mislaka Excel files (.xls and .xlsx) into structured data.
        
        Args:
            content: Raw Excel file bytes
            filename: Original filename
            ext: File extension ('xls' or 'xlsx')
            
        Returns:
            Dictionary with parsed data or None if parsing fails
        """
        try:
            rows = []
            columns = []
            
            if ext == 'xlsx':
                # Use openpyxl for .xlsx files
                try:
                    import openpyxl
                    from openpyxl import load_workbook
                    
                    wb = load_workbook(filename=io.BytesIO(content), data_only=True)
                    
                    for sheet_name in wb.sheetnames:
                        sheet = wb[sheet_name]
                        sheet_rows = list(sheet.iter_rows(values_only=True))
                        
                        if not sheet_rows:
                            continue
                        
                        # First row as headers
                        header_row = sheet_rows[0] if sheet_rows else []
                        sheet_columns = [str(h).strip() if h else f'עמודה_{i}' for i, h in enumerate(header_row)]
                        
                        # Add columns that don't exist yet
                        for col in sheet_columns:
                            if col and col not in columns:
                                columns.append(col)
                        
                        # Process data rows
                        for row in sheet_rows[1:]:
                            if row and any(cell is not None for cell in row):
                                row_dict = {}
                                for i, cell in enumerate(row):
                                    if i < len(sheet_columns):
                                        col_name = sheet_columns[i]
                                        row_dict[col_name] = cell if cell is not None else ''
                                rows.append(row_dict)
                    
                    wb.close()
                    
                except ImportError:
                    logger.warning("openpyxl not available for xlsx parsing")
                    return None
                    
            elif ext == 'xls':
                # Use xlrd for older .xls files
                try:
                    import xlrd
                    
                    wb = xlrd.open_workbook(file_contents=content)
                    
                    for sheet_idx in range(wb.nsheets):
                        sheet = wb.sheet_by_index(sheet_idx)
                        
                        if sheet.nrows == 0:
                            continue
                        
                        # First row as headers
                        header_row = sheet.row_values(0) if sheet.nrows > 0 else []
                        sheet_columns = [str(h).strip() if h else f'עמודה_{i}' for i, h in enumerate(header_row)]
                        
                        # Add columns that don't exist yet
                        for col in sheet_columns:
                            if col and col not in columns:
                                columns.append(col)
                        
                        # Process data rows
                        for row_idx in range(1, sheet.nrows):
                            row = sheet.row_values(row_idx)
                            if row and any(cell for cell in row):
                                row_dict = {}
                                for i, cell in enumerate(row):
                                    if i < len(sheet_columns):
                                        col_name = sheet_columns[i]
                                        row_dict[col_name] = cell if cell else ''
                                rows.append(row_dict)
                    
                except ImportError:
                    logger.warning("xlrd not available for xls parsing")
                    return None
            
            if not rows and not columns:
                return None
            
            # Map Excel columns to pension data structure
            return self._map_excel_to_pension_data(columns, rows, filename)
            
        except Exception as e:
            logger.error(f"Error parsing Excel file {filename}: {e}")
            return None
    
    def _parse_mislaka_csv(self, content: bytes, filename: str) -> Optional[Dict[str, Any]]:
        """
        Parse Mislaka CSV files into structured data.
        
        Args:
            content: Raw CSV file bytes
            filename: Original filename
            
        Returns:
            Dictionary with parsed data or None if parsing fails
        """
        try:
            # Try to decode with various encodings
            text_content = None
            for encoding in ['utf-8', 'windows-1255', 'iso-8859-8', 'cp1255']:
                try:
                    text_content = content.decode(encoding)
                    break
                except:
                    continue
            
            if not text_content:
                text_content = content.decode('utf-8', errors='replace')
            
            rows = []
            columns = []
            
            lines = text_content.strip().split('\n')
            if not lines:
                return None
            
            # Detect delimiter
            first_line = lines[0]
            delimiter = ','
            for delim in [',', '\t', ';', '|']:
                if delim in first_line:
                    delimiter = delim
                    break
            
            # Parse header
            columns = [col.strip().strip('"') for col in first_line.split(delimiter)]
            
            # Parse data rows
            for line in lines[1:]:
                if not line.strip():
                    continue
                values = [v.strip().strip('"') for v in line.split(delimiter)]
                if values and any(v for v in values):
                    row_dict = {}
                    for i, val in enumerate(values):
                        if i < len(columns):
                            row_dict[columns[i]] = val
                    rows.append(row_dict)
            
            if not rows:
                return None
            
            # Map CSV columns to pension data structure
            return self._map_excel_to_pension_data(columns, rows, filename)
            
        except Exception as e:
            logger.error(f"Error parsing CSV file {filename}: {e}")
            return None
    
    def _map_excel_to_pension_data(self, columns: List[str], rows: List[Dict], filename: str) -> Dict[str, Any]:
        """
        Map Excel/CSV column data to Mislaka pension data structure.
        
        Args:
            columns: List of column names
            rows: List of row dictionaries
            filename: Source filename
            
        Returns:
            Dictionary in standard Mislaka pension data format
        """
        if not looks_like_pension_table(columns):
            logger.info(f"File {filename} does not appear to contain pension data")
            return None
        
        mapped_columns = {}
        for col in columns:
            mapped = map_hebrew_column(col)
            if mapped:
                mapped_columns[col] = mapped
        
        # Extract client and account data
        client_info = {}
        accounts = []
        employers = []
        
        for row in rows:
            account = {
                'source': 'excel',
                'source_file': filename
            }
            
            for original_col, value in row.items():
                if value is None or str(value).strip() == '':
                    continue
                    
                if original_col in mapped_columns:
                    mapped_name = mapped_columns[original_col]
                    value_str = str(value).strip()
                    
                    # Convert value based on field type
                    if mapped_name in ['total_balance', 'savings_balance', 'severance_balance', 
                                      'management_fee', 'management_fee_savings', 'management_fee_deposits',
                                      'death_coverage', 'disability_coverage']:
                        try:
                            # Clean numeric value
                            clean_val = value_str.replace(',', '').replace('₪', '').replace('ש"ח', '').strip()
                            account[mapped_name] = float(clean_val)
                        except:
                            account[mapped_name] = 0
                    elif mapped_name == 'section14':
                        account[mapped_name] = value_str.lower() in ['כן', 'yes', '1', 'true', 'v', '✓', 'y']
                    elif mapped_name in ['full_name', 'first_name', 'last_name', 'id_number', 
                                        'birth_date', 'phone', 'mobile', 'email', 'address']:
                        # Client info
                        if not client_info.get(mapped_name):
                            client_info[mapped_name] = value_str
                    elif mapped_name == 'employer_name':
                        if value_str and value_str not in employers:
                            employers.append(value_str)
                        account[mapped_name] = value_str
                    else:
                        account[mapped_name] = value_str
            
            # Only add if we have some account data
            if account.get('provider') or account.get('policy_number') or account.get('total_balance'):
                accounts.append(account)
        
        # Build full name if we have parts
        if client_info.get('first_name') or client_info.get('last_name'):
            parts = [client_info.get('first_name', ''), client_info.get('last_name', '')]
            client_info['full_name'] = ' '.join(p for p in parts if p)
        
        # Translate product types
        for account in accounts:
            if account.get('product_type'):
                pt = account['product_type']
                for code, info in self.schema_mapping.PRODUCT_TYPE_CODES.items():
                    if pt in [info['he'], info['en'], code]:
                        account['product_type_name'] = info['he']
                        account['product_type_en'] = info['en']
                        break
            
            # Set default status if not present
            if not account.get('status'):
                account['status'] = 'פעיל'
                account['status_en'] = 'Active'
        
        return {
            'header': {
                'source': 'Excel/CSV',
                'filename': filename,
                'interface_type': 'Excel Import',
                'interface_type_he': 'יבוא מאקסל',
            },
            'client': client_info,
            'providers': [],  # Will be derived from accounts
            'accounts': accounts,
            'contributions': [],
            'severance': [],
            'employers': [{'name': e} for e in employers],
            'interface_type': 'Excel',
            'interface_type_he': 'יבוא מאקסל',
        }
    

    # ------------------------------------------------------------------------
    # Element parsers
    # ------------------------------------------------------------------------

    def _parse_header(self, root) -> Dict[str, Any]:
        """Parse header (KoteretKovetz) from XML."""
        header_elem = self._first_named(root, *self.HEADER_TAGS)
        if header_elem is None:
            header_elem = root
        return self._header_from_elem(header_elem)

    def _header_from_elem(self, header_elem) -> Dict[str, Any]:
        header = {}

        # Extract fields using mapping
        for xml_tag, field_name, variants in CompiledFields.HEADER:
            value = self._find_text(header_elem, xml_tag, variants)
            if value:
                if field_name == 'interface_code':
                    try:
                        header[field_name] = int(value)
                    except:
                        header[field_name] = 1
                else:
                    header[field_name] = value

        # Add interface type name
        interface_code = header.get('interface_code', 1)
        interface_info = self.schema_mapping.INTERFACE_CODES.get(interface_code, {})
        header['interface_type'] = interface_info.get('name', f'Type{interface_code}')
        header['interface_type_he'] = interface_info.get('he', f'סוג {interface_code}')

        # Format dates
        if header.get('created_at') and len(header['created_at']) == 14:
            try:
                dt = datetime.strptime(header['created_at'], '%Y%m%d%H%M%S')
                header['created_at_formatted'] = dt.strftime('%d/%m/%Y %H:%M')
            except:
                pass

        return header

    def _parse_client(self, root) -> Dict[str, Any]:
        """Parse client (YeshutLakoach) from XML."""
        client_elem = self._first_named(root, *self.CLIENT_TAGS)
        if client_elem is None:
            return {}
        return self._client_from_elem(client_elem)

    def _client_from_elem(self, client_elem) -> Dict[str, Any]:
        client = {}

        # Extract fields
        for xml_tag, field_name, variants in CompiledFields.CLIENT:
            value = self._find_text(client_elem, xml_tag, variants)
            if value:
                client[field_name] = value

        # Build full name if not present
        if not client.get('full_name') and (client.get('first_name') or client.get('last_name')):
            parts = [client.get('first_name', ''), client.get('last_name', '')]
            client['full_name'] = ' '.join(p for p in parts if p)

        # Translate ID type
        if client.get('id_type'):
            id_type_info = self.schema_mapping.ID_TYPE_CODES.get(client['id_type'], {})
            client['id_type_name'] = id_type_info.get('he', client['id_type'])

        return client

    def _parse_providers_and_accounts(self, root) -> Tuple[List[Dict], List[Dict]]:
        """Parse providers (YeshutYatzran) and accounts (HeshbonOPolisa)."""
        providers: List[Dict] = []
        accounts: List[Dict] = []

        for provider_elem in self._descendants_named(root, self.PROVIDER_TAG):
            self._parse_provider_elem(provider_elem, providers, accounts)

        for tag in self.STANDALONE_ACCOUNT_TAGS:
            for account_elem in self._descendants_named(root, tag):
                policy_num = self._find_text(account_elem, 'MISPAR-POLISA-O-HESHBON')
                if policy_num and not any(a.get('policy_number') == policy_num for a in accounts):
                    account = self._parse_account(account_elem, {}, {})
                    if account:
                        accounts.append(account)

        return providers, accounts

    def _parse_provider_elem(self, provider_elem, providers: List[Dict], accounts: List[Dict],
                             account_ids: Optional[set] = None) -> None:
        """One ``YeshutYatzran``: its fields, then every account under each of
        its products. ``account_ids`` (when given) collects ``id()`` of the
        account elements parsed here."""
        provider = {}

        for xml_tag, field_name, variants in CompiledFields.PROVIDER:
            value = self._find_text(provider_elem, xml_tag, variants)
            if value:
                provider[field_name] = value

        if provider:
            providers.append(provider)

        for product_elem in self._descendants_named(provider_elem, self.PRODUCT_TAG):
            product_info = {}

            for xml_tag, field_name, variants in CompiledFields.PRODUCT:
                value = self._find_text(product_elem, xml_tag, variants)
                if value:
                    product_info[field_name] = value

            for tag in self.ACCOUNT_TAGS:
                for account_elem in self._descendants_named(product_elem, tag):
                    account = self._parse_account(account_elem, provider, product_info)
                    if account:
                        accounts.append(account)
                        if account_ids is not None:
                            account_ids.add(id(account_elem))
            # The product's indexes are not needed again; keep the tree parse's
            # peak memory close to the document's own size.
            self._forget_index(product_elem)

    def _parse_account(self, elem, provider: Dict, product_info: Dict) -> Dict[str, Any]:
        """Parse a single account element."""
        account = {
            'provider': provider.get('name', ''),
            'provider_code': provider.get('code', ''),
            'product_type': product_info.get('product_type', ''),
            'product_type_code': product_info.get('product_type_code', ''),
            'product_name': product_info.get('name', ''),
        }

        # Extract account fields
        for xml_tag, field_name, variants in CompiledFields.ACCOUNT:
            value = self._find_text(elem, xml_tag, variants)
            if value:
                # Convert numeric fields
                if field_name in CompiledFields.ACCOUNT_NUMERIC:
                    account[field_name] = self._parse_number(value)
                elif field_name == 'section14':
                    account[field_name] = value in CompiledFields.TRUTHY
                else:
                    account[field_name] = value

        # Translate product type
        product_type_code = account.get('product_type_code', '') or account.get('product_type', '')
        if product_type_code in self.schema_mapping.PRODUCT_TYPE_CODES:
            type_info = self.schema_mapping.PRODUCT_TYPE_CODES[product_type_code]
            account['product_type_name'] = type_info['he']
            account['product_type_en'] = type_info['en']

        # Translate status
        status_code = account.get('status_code', '')
        if status_code in self.schema_mapping.STATUS_CODES:
            status_info = self.schema_mapping.STATUS_CODES[status_code]
            account['status'] = status_info['he']
            account['status_en'] = status_info['en']
        elif not account.get('status'):
            account['status'] = 'פעיל'
            account['status_en'] = 'Active'

        self._drop_component_total(elem, account)
        self._harvest_component_balances(elem, account)
        return account

    def _parse_contributions(self, root) -> List[Dict[str, Any]]:
        """Parse contributions (NetuneiHafrasha / PirteiHafrasha)."""
        contributions = []
        for tag in self.CONTRIBUTION_TAGS:
            for elem in self._descendants_named(root, tag):
                contrib = self._contribution_from_elem(elem)
                self._forget_index(elem)
                if contrib:
                    contributions.append(contrib)
        return contributions

    def _contribution_from_elem(self, elem) -> Dict[str, Any]:
        contrib = {}
        for xml_tag, field_name, variants in CompiledFields.CONTRIBUTION:
            value = self._find_text(elem, xml_tag, variants)
            if value:
                if field_name in CompiledFields.CONTRIBUTION_NUMERIC:
                    contrib[field_name] = self._parse_number(value)
                else:
                    contrib[field_name] = value
        return contrib

    def _parse_severance(self, root, interface_code: int) -> List[Dict[str, Any]]:
        """Parse severance data (NetuneiPitzuim)."""
        severance_list = []
        for tag in self.SEVERANCE_TAGS:
            for elem in self._descendants_named(root, tag):
                sev = self._severance_from_elem(elem)
                self._forget_index(elem)
                if sev:
                    severance_list.append(sev)
        return severance_list

    def _severance_from_elem(self, elem) -> Dict[str, Any]:
        sev = {}
        for xml_tag, field_name, variants in CompiledFields.SEVERANCE:
            value = self._find_text(elem, xml_tag, variants)
            if value:
                if field_name in CompiledFields.SEVERANCE_NUMERIC:
                    sev[field_name] = self._parse_number(value)
                elif field_name == 'section14':
                    # Section 14 codes: 1 = Yes, 2 = No (from schema)
                    sev[field_name] = value in CompiledFields.TRUTHY
                else:
                    sev[field_name] = value
        return sev

    def _parse_employers(self, root, accounts: List[Dict]) -> List[Dict[str, Any]]:
        """Parse employer information."""
        employers = []
        seen = set()

        # Collect from accounts
        for acct in accounts:
            emp_id = acct.get('employer_id', '')
            emp_name = acct.get('employer_name', '')
            if emp_name and emp_name not in seen:
                seen.add(emp_name)
                employers.append({'id': emp_id, 'name': emp_name})

        for elem in self._descendants_named(root, self.EMPLOYER_TAG):
            emp_name = self._find_text(elem, 'SHEM-MAASIK') or self._find_text(elem, 'ShemMaasik')
            emp_id = self._find_text(elem, 'KOD-MAASIK') or self._find_text(elem, 'KodMaasik')
            if emp_name and emp_name not in seen:
                seen.add(emp_name)
                employers.append({'id': emp_id, 'name': emp_name})

        return employers

    def _extract_all_elements(self, root) -> Dict[str, List[str]]:
        """Extract all text elements for additional analysis."""
        elements = {}

        def extract(elem, path=''):
            tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag

            if elem.text and elem.text.strip():
                if tag not in elements:
                    elements[tag] = []
                elements[tag].append(elem.text.strip())

            for child in elem:
                extract(child)

        extract(root)
        return elements

    def _find_text(self, elem, tag: str, variants: Optional[Tuple[str, ...]] = None) -> Optional[str]:
        """Find text content of a tag, trying multiple naming conventions.

        Inside a parse the lookup is a dict read on the element's descendant
        index; outside one it scans like ``find('.//tag')`` did. Both return
        the text of the *first* matching descendant per spelling (exact,
        hyphens removed, CamelCase), in that order of preference.
        """
        if elem is None:
            return None
        if variants is None:
            variants = tag_variants(tag)

        index = self._text_index(elem)
        if index is not None:
            for variant in variants:
                if variant in index:
                    text = index[variant]
                    if text:
                        return text.strip()
            return None

        for variant in variants:
            found = self._first_named(elem, variant)
            if found is not None and found.text:
                return found.text.strip()
        return None

    def _direct_text(self, elem, *tags) -> Optional[str]:
        """Prefer a direct child's text so nested Yitra blocks are not double-counted."""
        wanted = set(tags)
        for child in list(elem):
            if self._local_tag(child.tag) in wanted and child.text and child.text.strip():
                return child.text.strip()
        for tag in tags:
            value = self._find_text(elem, tag)
            if value:
                return value
        return None

    def _account_level_index(self, elem) -> Dict[str, Optional[str]]:
        """``tag → text of the first descendant with that tag``, skipping Yitra
        component subtrees so a component's amount is never read as the
        account's own figure."""
        index: Dict[str, Optional[str]] = {}
        stack = list(reversed(list(elem)))
        while stack:
            node = stack.pop()
            local = self._local_tag(node.tag)
            if local in self.YITRA_BLOCK_TAGS:
                continue
            if local not in index:
                index[local] = node.text
            stack.extend(reversed(list(node)))
        return index

    def _drop_component_total(self, elem, account: Dict[str, Any]) -> None:
        """Holdings aliases such as ``SACH-YITRA`` also name the amount inside a
        ``Yitra`` component, and ``_find_text`` takes the first matching
        descendant. Re-read the total outside the components so one component
        cannot stand in for the account's own holdings; without an
        account-level total the harvest below sums the components instead."""
        if 'total_balance' not in account:
            return
        if self._first_named(elem, *self.YITRA_BLOCK_TAGS) is None:
            return

        index = self._account_level_index(elem)
        total = None
        for variants in self.TOTAL_BALANCE_VARIANTS:
            for variant in variants:
                text = index.get(variant)
                if text and text.strip():
                    total = text.strip()
                    break
        if total is None:
            account.pop('total_balance')
        else:
            account['total_balance'] = self._parse_number(total)

    def _harvest_component_balances(self, elem, account: Dict[str, Any]) -> None:
        """Read official Yitra / PerutYitra children (KOD-SUG-HAFRASHA 1/2/3)."""
        savings = 0.0
        severance = 0.0
        found = False
        for yitra in self._descendants_named(elem, *self.YITRA_BLOCK_TAGS):
            if any(self._descendants_named(yitra, *self.YITRA_BLOCK_TAGS)):
                continue
            amount_text = self._direct_text(yitra, *self.YITRA_AMOUNT_TAGS)
            if not amount_text:
                continue
            amount = self._parse_number(amount_text)
            found = True
            code = (self._direct_text(yitra, *self.YITRA_TYPE_TAGS) or '').strip().lower()
            if code in {'3', '03', 'פיצויים', 'pitzuim', 'severance'}:
                severance += amount
            else:
                savings += amount
        if not found:
            return
        if not account.get('savings_balance'):
            account['savings_balance'] = savings
        if not account.get('severance_balance'):
            account['severance_balance'] = severance
        if not account.get('total_balance'):
            account['total_balance'] = savings + severance

    def _recover_client_identity(self, data: Dict[str, Any]) -> None:
        """Fill client.id_number from nested / raw affiliated tags when the header block omitted it."""
        client = data.get('client')
        if isinstance(client, list):
            client = client[0] if client else {}
        if not isinstance(client, dict):
            client = {}
        if not client.get('id_number'):
            raw = data.get('raw_elements') or {}
            for tag in self.CLIENT_ID_RAW_TAGS:
                values = raw.get(tag) or []
                if values:
                    client['id_number'] = str(values[0]).strip()
                    break
        if not client.get('full_name') and (client.get('first_name') or client.get('last_name')):
            client['full_name'] = ' '.join(
                p for p in (client.get('first_name', ''), client.get('last_name', '')) if p
            )
        data['client'] = client

    def _parse_number(self, value: str) -> float:
        """Parse numeric string to float, including Israeli ``1,000.50`` / ``1000,50`` forms."""
        if not value:
            return 0.0
        try:
            cleaned = str(value).strip().replace(' ', '').replace("'", '')
            cleaned = cleaned.replace('₪', '').replace('$', '').replace('€', '')
            cleaned = cleaned.replace('ש"ח', '').replace('ש״ח', '')
            if ',' in cleaned and '.' in cleaned:
                if cleaned.rfind(',') > cleaned.rfind('.'):
                    cleaned = cleaned.replace('.', '').replace(',', '.')
                else:
                    cleaned = cleaned.replace(',', '')
            elif ',' in cleaned and '.' not in cleaned:
                parts = cleaned.split(',')
                if len(parts) == 2 and 1 <= len(parts[1]) <= 2:
                    cleaned = parts[0] + '.' + parts[1]
                else:
                    cleaned = cleaned.replace(',', '')
            return float(cleaned)
        except Exception:
            return 0.0


class _ParseContext:
    """Per-parse scratch: descendant text indexes and stream positions, keyed
    by ``id(element)`` and discarded with the parse."""

    __slots__ = ('index', 'positions')

    def __init__(self):
        self.index: Dict[int, Dict[str, Optional[str]]] = {}
        self.positions: Dict[int, int] = {}


class _StreamAccumulator:
    """What the streaming parser collects before ``_assemble``."""

    CANDIDATE_TAGS = frozenset(
        MislakaParserMixin.HEADER_TAGS + MislakaParserMixin.CLIENT_TAGS
        + MislakaParserMixin.STANDALONE_ACCOUNT_TAGS + MislakaParserMixin.CONTRIBUTION_TAGS
        + MislakaParserMixin.SEVERANCE_TAGS + (MislakaParserMixin.EMPLOYER_TAG,)
    )

    def __init__(self):
        self.header_candidates: List[Tuple[int, float, Dict[str, Any]]] = []
        self.client_candidates: List[Tuple[int, float, Dict[str, Any]]] = []
        self.providers: List[Dict[str, Any]] = []
        self.provider_accounts: List[Dict[str, Any]] = []
        self.standalone: List[Tuple[int, float, str, Dict[str, Any]]] = []
        self.contributions: List[Tuple[int, float, Dict[str, Any]]] = []
        self.severance: List[Tuple[int, float, Dict[str, Any]]] = []
        self.employer_elems: List[Tuple[float, Optional[str], Optional[str]]] = []
        self.raw: Dict[str, List[str]] = {}


__all__ = ['MislakaParserMixin']
