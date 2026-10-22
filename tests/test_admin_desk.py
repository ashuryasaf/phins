"""Admin Desk registry, diary, documents, and the refusal to send."""

import base64
import csv
import hashlib
import io
import json
import os
import struct
import zipfile
from urllib.error import HTTPError
from urllib.request import Request, urlopen

import pytest

import services.admin_desk_service as desk


@pytest.fixture(autouse=True)
def _fresh_desk(monkeypatch, tmp_path):
    monkeypatch.delenv("PHINS_ADMIN_DESK_MAIL_ASAF", raising=False)
    monkeypatch.delenv("GODADDY_API_TOKEN", raising=False)
    monkeypatch.delenv("GODADDY_PAT", raising=False)
    monkeypatch.delenv("GODADDY_API_KEY", raising=False)
    monkeypatch.delenv("GODADDY_API_SECRET", raising=False)
    monkeypatch.setenv("PHINS_ADMIN_DESK_MAIL_VAULT_PATH", str(tmp_path / "mail-vault.json"))
    desk.reset_for_tests()
    yield
    desk.reset_for_tests()


def test_seed_mailbox_is_disabled_and_hides_the_secret():
    rows = desk.list_mailboxes("admin")
    admin = next(row for row in rows if row["id"] == "mbx-admin")
    assert admin["address"] == "asaf@phins.ai"
    assert admin["enabled"] is False
    assert admin["secret_configured"] is False
    assert admin["password_saved"] is False
    assert admin["secret_ref"] == "PHINS_ADMIN_DESK_MAIL_ASAF"
    assert "PHINS_ADMIN_DESK_MAIL_ASAF" not in json.dumps({k: v for k, v in admin.items() if k != "secret_ref"})
    stored = desk._read(desk._store(), "mbx-admin")
    assert desk.intact(stored)


def test_other_desks_are_slots_on_the_same_registry():
    rows = {row["desk"]: row for row in desk.list_mailboxes("admin")}
    assert set(rows) == {"admin", "actuary", "underwriting", "claims", "accountant", "compliance"}
    assert rows["actuary"]["address"] == ""
    assert rows["claims"]["signature_id"] == "sig-claims-v1"


def test_claims_session_cannot_see_the_admin_mailbox():
    visible = [row["id"] for row in desk.list_mailboxes("claims_adjuster")]
    assert visible == ["mbx-claims"]
    with pytest.raises(desk.DeskError) as raised:
        desk.get_mailbox("mbx-admin", "claims_adjuster")
    assert raised.value.status == 404


def test_enable_refuses_without_a_secret(monkeypatch):
    with pytest.raises(desk.DeskError) as raised:
        desk.set_mailbox_enabled("mbx-admin", True, "admin", "admin")
    assert raised.value.status == 409
    assert raised.value.message == "secret_not_configured"
    monkeypatch.setenv("PHINS_ADMIN_DESK_MAIL_ASAF", "desk-secret")
    enabled = desk.set_mailbox_enabled("mbx-admin", True, "admin", "admin")
    assert enabled["enabled"] is True
    assert enabled["secret_configured"] is True
    assert "desk-secret" not in json.dumps(enabled)


def test_actuary_cannot_enable_a_mailbox():
    with pytest.raises(desk.DeskError) as raised:
        desk.set_mailbox_enabled("mbx-actuary", True, "actuary", "actuary")
    assert raised.value.status == 403


def test_diary_import_is_idempotent_and_keeps_a_contradiction():
    first = desk.import_diary(
        [{"id": "diary-seed-01", "date": "2026-05-26", "time": "—", "counterparty": "Someone else",
          "location": "—", "purpose": "Different purpose", "status": "Held"}],
        "admin", "admin",
    )
    assert first["duplicate"] is False
    assert first["contradictions"] == 1
    assert first["added"] == 0
    again = desk.import_diary(
        [{"id": "diary-seed-01", "date": "2026-05-26", "time": "—", "counterparty": "Someone else",
          "location": "—", "purpose": "Different purpose", "status": "Held"}],
        "admin", "admin",
    )
    assert again["duplicate"] is True
    original = next(row for row in desk.list_diary("admin") if row["id"] == "diary-seed-01")
    assert original["counterparty"] == "Executive Actuary (IL · CA · US)"
    assert len(desk.list_contradictions("admin")) == 1


def test_letter_files_share_a_payload_and_keep_their_own_hashes():
    letter = desk.create_letter("Hello", "One line.\nTwo line.", "admin", "admin", "admin")
    assert letter["signature_id"] == "sig-admin-v1"
    html = desk.document_file(letter["id"], "html", "admin")
    pdf = desk.document_file(letter["id"], "pdf", "admin")
    docx = desk.document_file(letter["id"], "docx", "admin")
    html_bytes = base64.b64decode(html["data_b64"])
    assert b"PHINS" in html_bytes
    assert b"Ashury Asaf" in html_bytes
    assert hashlib.sha256(html_bytes).hexdigest() == html["sha256"]
    assert pdf["sha256"] != html["sha256"]
    package = zipfile.ZipFile(io.BytesIO(base64.b64decode(docx["data_b64"])))
    assert "Ashury Asaf" in package.read("word/document.xml").decode("utf-8")
    again = desk._letter_pdf("Hello", "One line.\nTwo line.", desk._footer_lines("admin", "asaf@phins.ai"))
    assert hashlib.sha256(again).hexdigest() == pdf["sha256"]


def test_table_csv_and_xlsx_carry_the_same_rows_and_require_a_source():
    with pytest.raises(desk.DeskError):
        desk.create_table("Queue", ["Item"], [["Diary"]], "", "admin", "admin", "admin")
    table = desk.create_table(
        "Queue", ["Item", "State"], [["Diary", "Ready"], ["Letter", "Ready"]],
        "Typed on the admin desk", "admin", "admin", "admin",
    )
    csv_file = desk.document_file(table["id"], "csv", "admin")
    xlsx_file = desk.document_file(table["id"], "xlsx", "admin")
    rows = list(csv.reader(io.StringIO(base64.b64decode(csv_file["data_b64"]).decode("utf-8"))))
    assert rows[0] == ["source", "Typed on the admin desk"]
    assert ["Diary", "Ready"] in rows
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(base64.b64decode(xlsx_file["data_b64"])))
    sheet_rows = [list(row) for row in book.active.iter_rows(values_only=True)]
    assert sheet_rows[0] == ("source", "Typed on the admin desk") or sheet_rows[0] == ["source", "Typed on the admin desk"]
    assert ("Diary", "Ready") in [tuple(row) for row in sheet_rows]


def test_send_is_refused_and_a_second_click_does_not_change_the_draft(monkeypatch):
    draft = desk.create_draft("mbx-admin", "person@example.com", "Hello", "Body", [], "admin", "admin")
    assert draft["from"] == "asaf@phins.ai"
    assert "Ashury Asaf" in draft["preview"]
    assert draft["sent"] is False
    first = desk.send_draft(draft["id"], "admin", "admin")
    assert first["status"] == "failed"
    assert first["send_error"] == "mailbox_not_ready"
    assert first["sent"] is False
    second = desk.send_draft(draft["id"], "admin", "admin")
    assert second["checksum"] == first["checksum"]
    monkeypatch.setenv("PHINS_ADMIN_DESK_MAIL_ASAF", "desk-secret")
    desk.set_mailbox_enabled("mbx-admin", True, "admin", "admin")
    fresh = desk.create_draft("mbx-admin", "person@example.com", "Hello", "Body", [], "admin", "admin")
    refused = desk.send_draft(fresh["id"], "admin", "admin")
    assert refused["send_error"] == "send_not_connected"
    assert refused["sent"] is False


def test_bad_checksum_is_not_served():
    store = desk._store()
    record = dict(desk._read(store, "mbx-admin"))
    record["address"] = "tampered@example.com"
    store["mbx-admin"] = record
    assert desk._read(store, "mbx-admin") is None
    assert desk.list_mailboxes("admin") == [] or all(row["id"] != "mbx-admin" for row in desk.list_mailboxes("admin"))
    faults = desk.overview("admin")["faults"]
    assert any(fault["subject_id"] == "mbx-admin" for fault in faults)


def _http(method, path, token=None, payload=None):
    base = os.environ["TEST_BASE_URL"]
    data = None if payload is None else json.dumps(payload).encode("utf-8")
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = "Bearer " + token
    request = Request(base + path, data=data, headers=headers, method=method)
    try:
        with urlopen(request, timeout=30) as response:
            return response.status, json.loads(response.read().decode("utf-8"))
    except HTTPError as exc:
        return exc.code, json.loads(exc.read().decode("utf-8"))


def _login(username, password):
    status, body = _http("POST", "/api/login", payload={"username": username, "password": password})
    assert status == 200, body
    return body["token"]


def _microsoft_dns(_domain):
    from services.admin_desk_mail import classify_dns
    return classify_dns(
        ["ns65.domaincontrol.com"],
        ["phins-ai.mail.protection.outlook.com"],
        ["v=spf1 include:secureserver.net -all", "NETORGFT20101992.onmicrosoft.com"],
    )


def _titan_dns(_domain):
    from services.admin_desk_mail import classify_dns
    return classify_dns(
        ["ns65.domaincontrol.com"],
        ["mailstore1.secureserver.net"],
        [],
    )


def test_dns_parser_reads_an_mx_answer():
    from services.admin_desk_mail import classify_dns, parse_dns
    qname = b"\x05phins\x02ai\x00"
    header = struct.pack(">HHHHHH", 0x1234, 0x8180, 1, 1, 0, 0)
    question = qname + struct.pack(">HH", 15, 1)
    exchange = b"\x08phins-ai\x04mail\x0aprotection\x07outlook\x03com\x00"
    rdata = b"\x00\x00" + exchange
    answer = struct.pack(">HHHIH", 0xC00C, 15, 1, 300, len(rdata)) + rdata
    parsed = parse_dns(header + question + answer, 0x1234)
    assert parsed["MX"] == ["phins-ai.mail.protection.outlook.com"]
    facts = classify_dns([], parsed["MX"], ["NETORGFT20101992.onmicrosoft.com"])
    assert facts["mail_product"] == "microsoft_365"
    assert facts["imap_host"] == "outlook.office365.com"
    assert facts["tenant_name"] == "NETORGFT20101992.onmicrosoft.com"


def test_godaddy_lists_the_admin_address_on_the_same_account(monkeypatch):
    from services import admin_desk_mail as mail
    monkeypatch.setenv("GODADDY_API_TOKEN", "super-secret-token")
    seen = []

    def transport(method, url, headers):
        seen.append(method)
        assert "super-secret-token" not in url
        assert headers["Authorization"] == "Bearer super-secret-token"
        if url.endswith("/v1/domains/phins.ai"):
            return 200, {"domain": "phins.ai"}
        if "check-mailbox-eligibility" in url:
            return 422, {"details": [{"issue": "EMAIL_ADDRESS_ALREADY_EXISTS"}]}
        if "/v1/email/mailboxes" in url:
            return 200, {"items": [{
                "emailAddress": "asaf@phins.ai",
                "mailboxType": "TITAN",
                "status": "COMPLETED",
            }]}
        raise AssertionError(url)

    report = mail.attribute_mailbox("asaf@phins.ai", transport=transport, dns_lookup=_microsoft_dns)
    assert seen and set(seen) == {"GET"}
    assert report["same_account"] is True
    assert report["verdict"] == "listed_on_account"
    assert report["same_account_basis"] == "mailbox_list"
    assert report["godaddy_api"]["mailbox_type"] == "TITAN"
    assert report["mail_product"] == "conflict"
    assert report["password_login"] is False
    assert "super-secret-token" not in json.dumps(report)


def test_missing_godaddy_token_still_reports_domain_dns():
    from services.admin_desk_mail import attribute_mailbox
    report = attribute_mailbox("asaf@phins.ai", dns_lookup=_microsoft_dns)
    assert report["verdict"] == "token_not_configured"
    assert report["same_account"] is None
    assert report["dns"]["nameservers_at_godaddy"] is True
    assert report["mail_product"] == "microsoft_365"
    assert report["dns"]["tenant_name"] == "NETORGFT20101992.onmicrosoft.com"
    assert report["dns"]["spf_includes_godaddy"] is True


def test_other_godaddy_account_is_not_the_same(monkeypatch):
    from services.admin_desk_mail import attribute_mailbox
    monkeypatch.setenv("GODADDY_API_TOKEN", "super-secret-token")

    def transport(method, url, headers):
        assert method == "GET"
        if "/v1/domains/" in url:
            return 404, {}
        if "/v1/email/mailboxes" in url:
            return 200, {"items": []}
        return 422, {"details": [{"issue": "EMAIL_PLAN_NOT_AVAILABLE"}]}

    report = attribute_mailbox("asaf@phins.ai", transport=transport, dns_lookup=_microsoft_dns)
    assert report["same_account"] is False
    assert report["verdict"] == "different_account"
    assert "super-secret-token" not in json.dumps(report)


def test_existing_address_on_an_owned_domain_counts_as_the_same_account(monkeypatch):
    from services.admin_desk_mail import attribute_mailbox
    monkeypatch.setenv("GODADDY_API_TOKEN", "super-secret-token")

    def transport(method, url, headers):
        if url.endswith("/v1/domains/phins.ai"):
            return 200, {"domain": "phins.ai"}
        if "/v1/email/mailboxes" in url:
            return 200, {"items": []}
        return 422, {"details": [{"issue": "EMAIL_ADDRESS_ALREADY_EXISTS"}]}

    report = attribute_mailbox("asaf@phins.ai", transport=transport, dns_lookup=_microsoft_dns)
    assert report["same_account"] is True
    assert report["verdict"] == "exists_on_owned_domain"
    assert report["mail_product"] == "microsoft_365"
    assert report["password_login"] is False


def test_password_is_saved_once_and_kept_out_of_the_desk_record(monkeypatch):
    from security.vault import VaultBlob
    calls = []

    def refuse_imap(*_args, **_kwargs):
        calls.append("imap")
        raise AssertionError("password must not be sent for Microsoft 365")

    saved = desk.save_mailbox_password(
        "mbx-admin", "desk-mail-password", "admin", "admin",
        dns_lookup=_microsoft_dns, imap_factory=refuse_imap,
    )
    assert calls == []
    assert saved["password_saved"] is True
    assert saved["secret_configured"] is True
    assert saved["connection_error"] == "microsoft_365_basic_auth_disabled"
    assert saved["connection_state"] == "saved"
    assert "desk-mail-password" not in json.dumps(saved)
    stored = json.dumps(desk._read(desk._store(), "mbx-admin"))
    assert "desk-mail-password" not in stored
    vault_text = open(os.environ["PHINS_ADMIN_DESK_MAIL_VAULT_PATH"], encoding="utf-8").read()
    assert "desk-mail-password" not in vault_text
    assert "fernet" in vault_text
    enabled = desk.set_mailbox_enabled("mbx-admin", True, "admin", "admin")
    assert enabled["enabled"] is True
    monkeypatch.setattr(
        "security.vault.encrypt_json",
        lambda _data: VaultBlob(scheme="plain", ciphertext="desk-mail-password"),
    )
    with pytest.raises(desk.DeskError) as raised:
        desk.save_mailbox_password("mbx-admin", "another-password", "admin", "admin")
    assert raised.value.status == 503
    assert "another-password" not in open(os.environ["PHINS_ADMIN_DESK_MAIL_VAULT_PATH"], encoding="utf-8").read()


def test_titan_mailbox_logs_in_once_with_the_saved_password():
    seen = {}

    class Client:
        def __init__(self, host, port, timeout=None):
            seen["host"] = host
            seen["port"] = port

        def login(self, user, password):
            seen["user"] = user
            seen["password"] = password
            return "OK", [b""]

        def logout(self):
            seen["logout"] = True
            return "OK", [b""]

    saved = desk.save_mailbox_password(
        "mbx-admin", "desk-mail-password", "admin", "admin",
        dns_lookup=_titan_dns, imap_factory=Client,
    )
    assert seen["host"] == "imap.secureserver.net"
    assert seen["port"] == 993
    assert seen["user"] == "asaf@phins.ai"
    assert seen["password"] == "desk-mail-password"
    assert seen["logout"] is True
    assert saved["connection_state"] == "connected"
    assert saved["connection_error"] == ""
    assert saved["connection_note"] == "Password saved, and the mailbox accepted it."
    assert "desk-mail-password" not in json.dumps(saved)


def test_refused_titan_login_stays_visible_on_the_mailbox():
    class Client:
        def __init__(self, host, port, timeout=None):
            pass

        def login(self, user, password):
            raise Exception("AUTHENTICATIONFAILED")

        def logout(self):
            return "OK", [b""]

    saved = desk.save_mailbox_password(
        "mbx-admin", "desk-mail-password", "admin", "admin",
        dns_lookup=_titan_dns, imap_factory=Client,
    )
    assert saved["connection_state"] == "saved"
    assert saved["connection_error"] == "auth_failed"
    assert saved["connection_note"] == "Password saved for email use. The mailbox refused the login."
    listed = next(row for row in desk.list_mailboxes("admin") if row["id"] == "mbx-admin")
    assert listed["connection_error"] == "auth_failed"
    assert listed["connection_note"] == saved["connection_note"]


def test_only_an_admin_can_save_a_mailbox_password():
    with pytest.raises(desk.DeskError) as raised:
        desk.save_mailbox_password("mbx-admin", "desk-mail-password", "actuary", "actuary")
    assert raised.value.status == 403
    with pytest.raises(desk.DeskError) as raised:
        desk.mailbox_attribution("mbx-admin", "claims_adjuster")
    assert raised.value.status == 403


def test_http_desk_scopes_mailboxes_and_serves_the_page():
    desk.reset_for_tests()
    admin = _login("admin", "admin123")
    page_request = Request(os.environ["TEST_BASE_URL"] + "/admin-desk.html")
    with urlopen(page_request, timeout=30) as response:
        page = response.read().decode("utf-8")
        assert response.status == 200
    assert "PHINS Admin Desk" in page
    assert "Save password" in page
    assert "auth_failed" in page and "probe_failed" in page
    status, body = _http("GET", "/api/admin/ai-agents/desk", token=admin)
    assert body["can_manage_mail"] is True
    assert status == 200
    assert any(row["address"] == "asaf@phins.ai" and row["enabled"] is False for row in body["mailboxes"])
    actuary = _login("actuary", "actuary123")
    status, body = _http("GET", "/api/admin/ai-agents/desk", token=actuary)
    assert status == 200
    assert [row["desk"] for row in body["mailboxes"]] == ["actuary"]
    status, body = _http("POST", "/api/admin/ai-agents/desk/drafts/missing/send", token=admin, payload={})
    assert status == 404
    status, report = _http("GET", "/api/admin/ai-agents/desk/mailboxes/mbx-admin/attribution", token=admin)
    assert status == 200, report
    assert report["address"] == "asaf@phins.ai"
    assert report["godaddy_api"]["reason"] == "godaddy_token_not_configured"
    status, saved = _http(
        "POST", "/api/admin/ai-agents/desk/mailboxes/mbx-admin/password",
        token=admin, payload={"password": "desk-mail-password"},
    )
    assert status == 200, saved
    assert saved["password_saved"] is True
    assert "desk-mail-password" not in json.dumps(saved)
    status, denied = _http(
        "POST", "/api/admin/ai-agents/desk/mailboxes/mbx-admin/password",
        token=actuary, payload={"password": "desk-mail-password"},
    )
    assert status == 403
    assert "desk-mail-password" not in json.dumps(denied)
